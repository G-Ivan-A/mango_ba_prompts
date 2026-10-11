#!/usr/bin/env python3
"""Транскрипция колонок 1-2 ТЗ BCREQ-1116 (issue #369) -> source_transcription.json.

Два исходника Заказчика — два блока отчёта:

* DOC «ТЕХНИЧЕСКОЕ ЗАДАНИЕ.doc» (Word 97-2003). Конвертируется LibreOffice в DOCX
  (`soffice --headless --convert-to docx`) только чтобы получить структуру абзацев,
  списков и таблиц; текст затем независимо сверяется с потоком WordDocument
  (`doc_text.py`, `check_source.py`).
  - Текст, удалённый Заказчиком в режиме исправлений (`w:del`, раздел
    «1. Существующая проблема и обоснование необходимости закупки»), в итоговой
    редакции документа отсутствует и не переносится.
  - Титул и «Глоссарий ключевых терминов и определений системы» (определения, не
    требования) исключены; пустые абзацы пропущены.
  - Номера «1.1.», «2.1.3.», «1.» … — литеральный текст абзаца, колонка 1 = номер,
    колонка 2 = остаток. Элементы маркированных списков Word (`numPr`, маркер
    Symbol U+F0B7, отображается «•») — колонка 1 = «•», колонка 2 = текст абзаца.
  - Строка таблицы — одна строка отчёта: колонка 1 пустая (в исходнике номера нет),
    колонка 2 = ячейки строки слева направо через `<br>`; продолжения вертикально
    объединённых ячеек (`vMerge` continue) пропускаются, знак сноски Word выводится
    как `<sup>1</sup>`.
* XLSX «Требования к ПО КЦ (4).xlsx», лист «функционал». Нумерации требований нет —
  колонка 1 = номер строки листа Excel, колонка 2 = значение ячейки столбца A
  (полужирные ячейки — заголовки групп). Колонка B («Ответ» / «да (если есть)/ нет
  (если отсутствует)») — графа ответа участника, требованием не является.

Текст переносится как есть, включая хвостовые пробелы и неразрывные пробелы.

    python3 experiments/issue_369/make_transcription.py <ТЕХНИЧЕСКОЕ ЗАДАНИЕ.doc> <Требования к ПО КЦ (4).xlsx>
"""
from __future__ import annotations

import hashlib
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
import openpyxl
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
DOC_TITLE = "ТЕХНИЧЕСКОЕ ЗАДАНИЕ.doc"
XLSX_TITLE = "Требования к ПО КЦ (4).xlsx"
FOOTNOTE = "<sup>1</sup>"


def run_text(el) -> str:
    """Текст абзаца без удалённых исправлений (`w:delText` не учитывается).

    `w:tab`, которые LibreOffice добавляет при конвертации (в потоке WordDocument
    символов табуляции нет), не переносятся.
    """
    return "".join(FOOTNOTE if node.tag == qn("w:footnoteReference") else node.text or ""
                   for node in el.iter(qn("w:t"), qn("w:footnoteReference")))


def doc_rows(path: Path) -> list[list]:
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--convert-to", "docx", "--outdir", tmp, str(path)],
                       check=True, capture_output=True)
        document = docx.Document(str(next(Path(tmp).glob("*.docx"))))
    body = list(document.element.body.iterchildren())
    texts = [run_text(el) if el.tag == qn("w:p") else None for el in body]
    start = texts.index("1. Назначение и цели проекта")
    assert texts[2] == "Глоссарий ключевых терминов и определений системы" and body[3].tag == qn("w:tbl")
    rows: list[list] = []
    tables = 0
    for idx, el in enumerate(body):
        if el.tag == qn("w:tbl"):
            tables += 1
            if idx < start:
                continue
            for rno, tr in enumerate(el.findall(qn("w:tr"))):
                cells = []
                for tc in tr.findall(qn("w:tc")):
                    vmerge = tc.find(f"{qn('w:tcPr')}/{qn('w:vMerge')}")
                    if vmerge is not None and vmerge.get(qn("w:val")) in (None, "continue"):
                        continue
                    cells.append("<br>".join(run_text(p) for p in tc.findall(qn("w:p"))))
                if cells:
                    rows.append(["", "<br>".join(cells), f"t{tables}r{rno + 1}"])
            continue
        if idx < start or el.tag != qn("w:p"):
            continue
        text = texts[idx]
        if not text.strip():
            continue
        if el.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None:
            rows.append(["•", text, f"p{idx}"])
            continue
        m = re.match(r"(\d+(?:\.\d+)*\.) (.*)\Z", text, re.S)
        rows.append([m.group(1), m.group(2), f"p{idx}"] if m else ["", text, f"p{idx}"])
    return [[f"D{i:03d}", *row] for i, row in enumerate(rows, start=1)]


def xlsx_rows(path: Path) -> list[list]:
    ws = openpyxl.load_workbook(str(path)).worksheets[0]
    assert ws.title == "функционал" and ws["B1"].value == "Ответ"
    rows = []
    for r in range(2, ws.max_row + 1):
        value = ws.cell(r, 1).value
        if value is None or not str(value).strip():
            continue
        assert ws.cell(r, 1).data_type == "s", r
        rows.append([str(r), value, f"A{r}" + ("|b" if ws.cell(r, 1).font.b else "")])
    return [[f"X{i:03d}", *row] for i, row in enumerate(rows, start=1)]


def block(key: str, title: str, path: Path, rows: list[list]) -> dict:
    raw = path.read_bytes()
    return {"key": key, "title": title, "sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw), "rows": rows}


if __name__ == "__main__":
    doc_path, xlsx_path = Path(sys.argv[1]), Path(sys.argv[2])
    data = {"blocks": [block("DOC", DOC_TITLE, doc_path, doc_rows(doc_path)),
                       block("XLSX", XLSX_TITLE, xlsx_path, xlsx_rows(xlsx_path))]}
    HERE.joinpath("source_transcription.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(", ".join(f"{b['key']}: {len(b['rows'])} rows" for b in data["blocks"]))
